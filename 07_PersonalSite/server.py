from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from io import StringIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote


ROOT = Path(__file__).resolve().parents[1]
APP_ROOT = Path(__file__).resolve().parent
WEB_ROOT = APP_ROOT / "web"
DATA_FILE = APP_ROOT / "data" / "showcase.json"
RAW_MEDIA_ROOT = APP_ROOT / "media" / "raw"
ROUTINE_ROOT = APP_ROOT / "local" / "routine"
ROUTINE_IMPORTS_ROOT = ROUTINE_ROOT / "imports"
ROUTINE_RECORDS_FILE = ROUTINE_ROOT / "routine-records.json"
MARGINALIA_RUNTIME = ROOT / ".marginalia-yusu"
MARGINALIA_DIST = APP_ROOT / "marginalia-dist"
MARGINALIA_BACKEND = APP_ROOT / "marginalia-backend"
KAOYAN_WORKSPACE = Path(
    os.environ.get(
        "YUSU_KAOYAN_WORKSPACE",
        str(ROOT.parent / "000资料相关" / "000考研"),
    )
)
KAOYAN_DASHBOARD = KAOYAN_WORKSPACE / "00_打开-北交电气考研数据看板.html"
KAOYAN_ALLOWED_ROOTS = [
    KAOYAN_WORKSPACE / "北京交通大学资料",
]
KAOYAN_ALLOWED_FILES = {
    KAOYAN_DASHBOARD,
    KAOYAN_WORKSPACE / "00_SUPERYUSU考研项目总览.md",
}

SEARCH_ROOTS = [
    ROOT / "01_Projects",
    ROOT / "02_GlobalMemory",
    ROOT / "03_CrossProject",
    ROOT / "04_Runbooks",
    ROOT / "06_Maps",
    ROOT / "07_PersonalSite",
    ROOT / "INDEX.md",
    ROOT / "README.md",
]


def _load_env_file(path: Path) -> None:
    if not path.is_file():
        return
    for raw_line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip("\"'")
        if key:
            os.environ.setdefault(key, value)


# Marginalia reads configuration while its modules are imported. Load the
# ignored repo-local runtime first, then enable its worker in this one process.
_load_env_file(MARGINALIA_RUNTIME / ".env")
os.environ.setdefault("MARGINALIA_HOME", str(MARGINALIA_RUNTIME / "data"))
os.environ["WORKER_ENABLED"] = os.environ.get("YUSU_MARGINALIA_WORKER", "true")
os.environ.setdefault("MARGINALIA_DESKTOP", "1")

if MARGINALIA_BACKEND.is_dir():
    sys.path.insert(0, str(MARGINALIA_BACKEND))

from fastapi import File, Form, HTTPException, Query, UploadFile  # noqa: E402
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse  # noqa: E402
from fastapi.staticfiles import StaticFiles  # noqa: E402
from marginalia.config import get_settings, resolve_profile  # noqa: E402
from marginalia.main import app  # noqa: E402


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="replace")


def _iter_markdown_files() -> list[Path]:
    files: list[Path] = []
    for root in SEARCH_ROOTS:
        if root.is_file() and root.suffix.lower() == ".md":
            files.append(root)
        elif root.is_dir():
            files.extend(path for path in root.rglob("*.md") if path.is_file())
    return sorted(files)


def _resolve_markdown_doc(rel_path: str) -> Path | None:
    clean = unquote(rel_path).replace("\\", "/").strip()
    if not clean or clean.startswith("/") or "\x00" in clean:
        return None
    if any(part == ".." for part in clean.split("/")):
        return None

    target = (ROOT / clean).resolve()
    if target.suffix.lower() != ".md":
        return None
    allowed = {path.resolve() for path in _iter_markdown_files()}
    return target if target in allowed else None


def _resolve_kaoyan_asset(asset_path: str) -> Path | None:
    clean = unquote(asset_path).replace("\\", "/").lstrip("/")
    if not clean or "\x00" in clean:
        return KAOYAN_DASHBOARD if KAOYAN_DASHBOARD.is_file() else None
    if any(part == ".." for part in clean.split("/")):
        return None

    target = (KAOYAN_WORKSPACE / clean).resolve()
    if not target.is_file():
        return None

    allowed_files = {path.resolve() for path in KAOYAN_ALLOWED_FILES}
    if target in allowed_files:
        return target

    for allowed_root in KAOYAN_ALLOWED_ROOTS:
        root = allowed_root.resolve()
        if root.exists() and _is_relative_to(target, root):
            return target
    return None


def _file_timestamp(path: Path) -> str | None:
    if not path.is_file():
        return None
    from datetime import datetime

    return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")


def _inject_kaoyan_return_button(html: str) -> str:
    if '<a data-yusu-return-link' in html:
        return html
    control = """
  <style>
    .yusu-return-link {
      position: fixed;
      top: 18px;
      right: 18px;
      z-index: 9999;
      display: inline-flex;
      align-items: center;
      min-height: 38px;
      padding: 0 14px;
      border: 1px solid rgba(15, 107, 99, 0.28);
      border-radius: 6px;
      background: rgba(255, 255, 255, 0.92);
      color: #0f2b2a;
      box-shadow: 0 12px 30px rgba(19, 30, 37, 0.14);
      backdrop-filter: blur(12px);
      font: 700 14px/1 "Microsoft YaHei", "Segoe UI", system-ui, sans-serif;
      text-decoration: none;
    }
    .yusu-return-link:hover {
      border-color: rgba(15, 107, 99, 0.5);
      color: #0f6b63;
      transform: translateY(-1px);
    }
    @media (max-width: 720px) {
      .yusu-return-link {
        top: 12px;
        right: 12px;
        min-height: 34px;
        padding: 0 10px;
        font-size: 12px;
      }
    }
  </style>
  <a data-yusu-return-link class="yusu-return-link" href="/" aria-label="返回 YUSU 主页">← 返回 YUSU</a>
"""
    return html.replace("</body>", f"{control}\n</body>", 1)


def _first_heading(text: str, fallback: str) -> str:
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("#"):
            return line.lstrip("#").strip() or fallback
    return fallback


def _snippet(lines: list[str], line_index: int, width: int = 1) -> str:
    start = max(0, line_index - width)
    end = min(len(lines), line_index + width + 1)
    joined = " ".join(part.strip() for part in lines[start:end] if part.strip())
    return joined[:420]


def search_markdown(query: str, limit: int = 14) -> list[dict]:
    terms = [term.lower() for term in re.split(r"\s+", query.strip()) if term.strip()]
    if not terms:
        return []

    results: list[dict] = []
    for path in _iter_markdown_files():
        try:
            text = _read_text(path)
        except OSError:
            continue
        lowered = text.lower()
        matched_terms = [term for term in terms if term in lowered]
        if not matched_terms:
            continue

        lines = text.splitlines()
        hit_index = next(
            (
                index
                for index, line in enumerate(lines)
                if any(term in line.lower() for term in matched_terms)
            ),
            0,
        )
        rel = path.relative_to(ROOT).as_posix()
        raw_score = sum(lowered.count(term) for term in matched_terms)
        results.append(
            {
                "title": _first_heading(text, path.stem),
                "path": rel,
                "line": hit_index + 1,
                "snippet": _snippet(lines, hit_index),
                "score": len(matched_terms) * 1000 + raw_score,
            }
        )

    results.sort(key=lambda item: (-item["score"], item["path"]))
    return results[:limit]


def _safe_filename(name: str) -> str:
    clean = unquote(name or "tomatodo-export").replace("\\", "_").replace("/", "_")
    clean = re.sub(r"[\x00-\x1f<>:\"|?*]+", "_", clean).strip(" ._")
    return clean[:120] or "tomatodo-export"


def _decode_table_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8-sig", errors="replace")


def _read_csv_like(path: Path, delimiter: str | None = None) -> list[list[Any]]:
    text = _decode_table_bytes(path.read_bytes())
    sample = text[:2048]
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = "\t" if "\t" in sample else ","
    return [row for row in csv.reader(StringIO(text), delimiter=delimiter)]


def _read_xlsx(path: Path) -> list[list[Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]


def _read_xls_with_xlrd(path: Path) -> list[list[Any]]:
    import xlrd

    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    rows: list[list[Any]] = []
    for row_index in range(sheet.nrows):
        values: list[Any] = []
        for col_index in range(sheet.ncols):
            cell = sheet.cell(row_index, col_index)
            if cell.ctype == xlrd.XL_CELL_DATE:
                values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
            else:
                values.append(cell.value)
        rows.append(values)
    return rows


def _read_xls_via_excel_com(path: Path) -> list[list[Any]]:
    script = r"""
$ErrorActionPreference = "Stop"
$path = $env:YUSU_ROUTINE_XLS_PATH
$out = $env:YUSU_ROUTINE_XLS_JSON
$excel = $null
$workbook = $null
try {
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $workbook = $excel.Workbooks.Open($path, 0, $true)
  $sheet = $workbook.Worksheets.Item(1)
  $range = $sheet.UsedRange
  $rows = New-Object System.Collections.Generic.List[object]
  for ($r = 1; $r -le $range.Rows.Count; $r++) {
    $values = New-Object System.Collections.Generic.List[string]
    for ($c = 1; $c -le $range.Columns.Count; $c++) {
      [void]$values.Add([string]$range.Cells.Item($r, $c).Text)
    }
    [void]$rows.Add($values.ToArray())
  }
  $json = ConvertTo-Json -InputObject $rows.ToArray() -Depth 5 -Compress
  [System.IO.File]::WriteAllText($out, $json, [System.Text.UTF8Encoding]::new($false))
}
finally {
  if ($workbook) { $workbook.Close($false) | Out-Null }
  if ($excel) { $excel.Quit() | Out-Null }
}
"""
    env = os.environ.copy()
    env["YUSU_ROUTINE_XLS_PATH"] = str(path)
    tmp_root = ROOT / ".tools" / "tmp"
    tmp_root.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, dir=tmp_root, encoding="utf-8") as handle:
        json_path = Path(handle.name)
    env["YUSU_ROUTINE_XLS_JSON"] = str(json_path)
    try:
        completed = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
            check=False,
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            env=env,
            timeout=60,
        )
        if completed.returncode != 0:
            detail = completed.stderr.strip() or completed.stdout.strip() or "Excel COM failed"
            raise HTTPException(status_code=422, detail=f"无法读取 .xls：{detail}")
        try:
            rows = json.loads(json_path.read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=422, detail=f"无法解析 .xls 输出：{exc}") from exc
    finally:
        json_path.unlink(missing_ok=True)
    return rows if isinstance(rows, list) else []


def _read_routine_table(path: Path) -> list[list[Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        if os.environ.get("YUSU_ALLOW_SERVER_XLS_IMPORT") == "1":
            try:
                return _read_xls_with_xlrd(path)
            except ImportError:
                return _read_xls_via_excel_com(path)
        raise HTTPException(
            status_code=422,
            detail=".xls 请通过 /routine/ 网页上传入口导入；浏览器端 SheetJS 解析可避免 Windows 后端直读老 xls 的中文编码问题。",
        )
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".csv":
        return _read_csv_like(path)
    if suffix == ".tsv":
        return _read_csv_like(path, delimiter="\t")
    raise HTTPException(status_code=415, detail="仅支持 .xls, .xlsx, .csv, .tsv")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return str(value).strip()


def _parse_tomatodo_time(value: Any) -> tuple[datetime, datetime] | None:
    text = _cell_text(value)
    match = re.search(
        r"(\d{4}-\d{1,2}-\d{1,2}\s+\d{1,2}:\d{2})\s*(?:至|~|-)\s*(\d{4}-\d{1,2}-\d{1,2}\s+\d{1,2}:\d{2})",
        text,
    )
    if not match:
        return None
    start = datetime.strptime(match.group(1), "%Y-%m-%d %H:%M")
    end = datetime.strptime(match.group(2), "%Y-%m-%d %H:%M")
    return start, end


def _parse_minutes(value: Any) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return max(0, int(round(float(value))))
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    return max(0, int(round(float(match.group(0))))) if match else 0


def _parse_completion(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return round(number * 100, 2) if 0 <= number <= 1 else round(number, 2)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    return round(float(match.group(0)), 2) if match else None


def _record_id(start: datetime, end: datetime, task: str, minutes: int, note: str) -> str:
    raw = f"{start.isoformat()}|{end.isoformat()}|{task}|{minutes}|{note}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:20]


def _parse_routine_records(rows: list[list[Any]], source_name: str) -> dict:
    header_index = None
    header: list[str] = []
    for index, row in enumerate(rows):
        labels = [_cell_text(cell) for cell in row]
        if "专注时间" in labels and "待办名称" in labels:
            header_index = index
            header = labels
            break
    if header_index is None:
        raise HTTPException(status_code=422, detail="没有找到番茄 ToDo 导出表头：专注时间 / 待办名称")

    def col(name: str) -> int | None:
        return header.index(name) if name in header else None

    time_col = col("专注时间")
    task_col = col("待办名称")
    minutes_col = col("专注时长(分钟)")
    note_col = col("心得")
    status_col = col("状态")
    completion_col = col("完成度")
    records: list[dict] = []

    for raw_row in rows[header_index + 1 :]:
        row = list(raw_row) + [""] * max(0, len(header) - len(raw_row))
        parsed_time = _parse_tomatodo_time(row[time_col] if time_col is not None else "")
        task = _cell_text(row[task_col] if task_col is not None else "")
        if not parsed_time or not task:
            continue
        start, end = parsed_time
        minutes = _parse_minutes(row[minutes_col] if minutes_col is not None else "")
        if minutes == 0 and end > start:
            minutes = max(0, int(round((end - start).total_seconds() / 60)))
        note = _cell_text(row[note_col] if note_col is not None else "")
        status = _cell_text(row[status_col] if status_col is not None else "")
        completion = _parse_completion(row[completion_col] if completion_col is not None else "")
        include = minutes > 0 and status != "中途放弃"
        records.append(
            {
                "id": _record_id(start, end, task, minutes, note),
                "start": start.isoformat(timespec="minutes"),
                "end": end.isoformat(timespec="minutes"),
                "date": start.date().isoformat(),
                "task": task,
                "minutes": minutes,
                "note": note,
                "status": status,
                "completion": completion,
                "includeInStats": include,
                "source": source_name,
            }
        )
    return {"header": header, "records": records}


def _normalize_client_routine_records(items: list[Any], source_name: str) -> list[dict]:
    records: list[dict] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        try:
            start = datetime.fromisoformat(str(item.get("start", "")).replace(" ", "T"))
            end = datetime.fromisoformat(str(item.get("end", "")).replace(" ", "T"))
        except ValueError:
            continue
        task = _cell_text(item.get("task"))
        if not task:
            continue
        minutes = _parse_minutes(item.get("minutes"))
        if minutes == 0 and end > start:
            minutes = max(0, int(round((end - start).total_seconds() / 60)))
        note = _cell_text(item.get("note"))
        status = _cell_text(item.get("status"))
        completion = _parse_completion(item.get("completion"))
        records.append(
            {
                "id": _record_id(start, end, task, minutes, note),
                "start": start.isoformat(timespec="minutes"),
                "end": end.isoformat(timespec="minutes"),
                "date": start.date().isoformat(),
                "task": task,
                "minutes": minutes,
                "note": note,
                "status": status,
                "completion": completion,
                "includeInStats": minutes > 0 and status != "中途放弃",
                "source": source_name,
            }
        )
    return records


def _load_routine_store() -> dict:
    if not ROUTINE_RECORDS_FILE.is_file():
        return {"version": 1, "updatedAt": None, "records": [], "imports": []}
    try:
        return json.loads(ROUTINE_RECORDS_FILE.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {"version": 1, "updatedAt": None, "records": [], "imports": []}


def _save_routine_store(store: dict) -> None:
    ROUTINE_ROOT.mkdir(parents=True, exist_ok=True)
    store["updatedAt"] = datetime.now().isoformat(timespec="seconds")
    temp = ROUTINE_RECORDS_FILE.with_suffix(".tmp")
    temp.write_text(json.dumps(store, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(ROUTINE_RECORDS_FILE)


def _merge_routine_records(new_records: list[dict], import_info: dict) -> dict:
    store = _load_routine_store()
    by_id = {record["id"]: record for record in store.get("records", []) if record.get("id")}
    added = 0
    updated = 0
    for record in new_records:
        existing = by_id.get(record["id"])
        if existing:
            sources = set(existing.get("sources", []))
            sources.add(record["source"])
            existing.update(record)
            existing["sources"] = sorted(sources)
            updated += 1
        else:
            record["sources"] = [record["source"]]
            by_id[record["id"]] = record
            added += 1
    imports = store.get("imports", [])
    imports.append(import_info | {"recordsSeen": len(new_records), "recordsAdded": added})
    store["imports"] = imports[-30:]
    store["records"] = sorted(by_id.values(), key=lambda item: (item["start"], item["id"]))
    store.pop("milestones", None)
    store.pop("milestonesGeneratedAt", None)
    store.pop("milestonesGeneratedBy", None)
    store.pop("milestoneSignature", None)
    _save_routine_store(store)
    return {"added": added, "updated": updated, "total": len(store["records"])}


def _iter_time_segments(start: datetime, end: datetime, minutes: int, mode: str) -> list[tuple[str, float]]:
    if minutes <= 0 or end <= start:
        return []
    segments: list[tuple[str, float]] = []
    cursor = start
    while cursor < end:
        if mode == "day":
            boundary = datetime.combine(cursor.date() + timedelta(days=1), datetime.min.time())
            key = cursor.date().isoformat()
        else:
            boundary = (cursor.replace(minute=0, second=0, microsecond=0) + timedelta(hours=1))
            key = f"{cursor.hour:02d}:00"
        next_cursor = min(end, boundary)
        raw_minutes = max(0.0, (next_cursor - cursor).total_seconds() / 60)
        if raw_minutes:
            segments.append((key, raw_minutes))
        cursor = next_cursor
    raw_total = sum(value for _, value in segments)
    if not raw_total:
        return []
    factor = minutes / raw_total
    return [(key, value * factor) for key, value in segments]


def _date_span(start: datetime, end: datetime) -> list[str]:
    days = []
    cursor = start.date()
    while cursor <= end.date():
        days.append(cursor.isoformat())
        cursor += timedelta(days=1)
    return days


def _month_span(start: datetime, end: datetime) -> list[datetime]:
    months: list[datetime] = []
    cursor = start.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    last = end.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    while cursor <= last:
        months.append(cursor)
        year = cursor.year + (1 if cursor.month == 12 else 0)
        month = 1 if cursor.month == 12 else cursor.month + 1
        cursor = cursor.replace(year=year, month=month)
    return months


def _routine_note_signature(records: list[dict]) -> str:
    note_rows = [
        f"{record.get('id', '')}|{record.get('date', '')}|{record.get('task', '')}|{record.get('note', '')}"
        for record in records
        if record.get("note")
    ]
    raw = "\n".join(sorted(note_rows))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:20]


def _short_text(text: str, limit: int = 28) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    return compact if len(compact) <= limit else compact[: limit - 1] + "…"


def _fallback_routine_milestones(records: list[dict]) -> list[dict]:
    milestones: list[dict] = []
    keyword_rules = [
        ("start", ("开始", "正式开始", "开了")),
        ("finish", ("看完", "结束", "完成", "整理完")),
        ("chapter", ("第一讲", "第二讲", "第三讲", "第四讲", "第五讲", "第六讲", "章节")),
        ("exercise", ("做题", "习题", "题")),
        ("accuracy", ("正确率", "准确率")),
        ("practice", ("习题", "正确率", "错题")),
        ("review", ("整理", "复盘", "总结")),
    ]
    for record in records:
        note = _cell_text(record.get("note"))
        if not note:
            continue
        milestone_type = "note"
        for candidate, keywords in keyword_rules:
            if any(keyword in note for keyword in keywords):
                milestone_type = candidate
                break
        metric = {}
        accuracy = re.search(r"(?:正确率|准确率)\s*[:：]?\s*(\d+(?:\.\d+)?)\s*%", note)
        if accuracy:
            metric["accuracy"] = f"{accuracy.group(1)}%"
            milestone_type = "accuracy"
        question_count = re.search(r"(\d+)\s*道题", note)
        if question_count:
            metric["questions"] = f"{question_count.group(1)}道"
        milestones.append(
            {
                "date": record.get("date"),
                "task": record.get("task"),
                "type": milestone_type,
                "label": metric.get("accuracy") or _short_text(note, 18),
                "detail": note,
                "metric": metric,
                "pinned": milestone_type in {"start", "finish", "chapter", "accuracy"},
                "recordId": record.get("id"),
                "confidence": 0.55,
                "source": "heuristic",
            }
        )
    return milestones


def _validate_routine_milestones(items: Any, records: list[dict]) -> list[dict]:
    valid_dates = {record.get("date") for record in records}
    valid_ids = {record.get("id") for record in records}
    allowed_types = {
        "start",
        "finish",
        "chapter",
        "exercise",
        "accuracy",
        "practice",
        "review",
        "warning",
        "note",
        "other",
    }
    milestones: list[dict] = []
    if not isinstance(items, list):
        return milestones
    for item in items:
        if not isinstance(item, dict):
            continue
        date = _cell_text(item.get("date"))
        if date not in valid_dates:
            continue
        label = _short_text(_cell_text(item.get("label")), 24)
        detail = _cell_text(item.get("detail"))
        if not label and detail:
            label = _short_text(detail, 24)
        if not label:
            continue
        record_id = _cell_text(item.get("recordId"))
        milestone_type = _cell_text(item.get("type")) or "other"
        if milestone_type not in allowed_types:
            milestone_type = "other"
        confidence = item.get("confidence", 0.7)
        try:
            confidence_value = max(0.0, min(1.0, float(confidence)))
        except (TypeError, ValueError):
            confidence_value = 0.7
        metric = item.get("metric") if isinstance(item.get("metric"), dict) else {}
        clean_metric = {str(key): _short_text(str(value), 24) for key, value in metric.items() if value not in (None, "")}
        milestones.append(
            {
                "date": date,
                "task": _short_text(_cell_text(item.get("task")), 28),
                "type": milestone_type,
                "label": label,
                "detail": detail or label,
                "metric": clean_metric,
                "pinned": bool(item.get("pinned", milestone_type in {"start", "finish", "chapter", "accuracy"})),
                "recordId": record_id if record_id in valid_ids else None,
                "confidence": round(confidence_value, 2),
                "source": _cell_text(item.get("source")) or "deepseek",
            }
        )
    milestones.sort(key=lambda item: (item["date"], item["type"], item["label"]))
    return milestones[:80]


def _milestones_for_payload(store: dict, records: list[dict]) -> tuple[list[dict], str]:
    signature = _routine_note_signature(records)
    stored = store.get("milestones")
    if store.get("milestoneSignature") == signature and isinstance(stored, list):
        return _validate_routine_milestones(stored, records), store.get("milestonesGeneratedBy") or "stored"
    return _fallback_routine_milestones(records), "heuristic"


def _build_routine_calendar(records: list[dict], milestones: list[dict]) -> list[dict]:
    dated_records = [record for record in records if record.get("date")]
    if not dated_records:
        return []
    starts = [datetime.fromisoformat(record["start"]) for record in dated_records if record.get("start")]
    ends = [datetime.fromisoformat(record["end"]) for record in dated_records if record.get("end")]
    if not starts or not ends:
        return []

    by_day: dict[str, dict] = {}
    for record in dated_records:
        date = record["date"]
        day = by_day.setdefault(
            date,
            {
                "date": date,
                "minutes": 0,
                "hours": 0,
                "tasks": {},
                "notes": [],
                "records": [],
                "milestones": [],
                "taskMinutes": {},
            },
        )
        if record.get("includeInStats"):
            minutes = int(record.get("minutes") or 0)
            day["minutes"] += minutes
            task = record.get("task") or "未命名"
            day["tasks"][task] = day["tasks"].get(task, 0) + minutes
            day["taskMinutes"][task] = day["taskMinutes"].get(task, 0) + minutes
        if record.get("note"):
            day["notes"].append(
                {
                    "recordId": record.get("id"),
                    "task": record.get("task"),
                    "note": record.get("note"),
                    "minutes": record.get("minutes", 0),
                }
            )
        day["records"].append(
            {
                "id": record.get("id"),
                "start": record.get("start"),
                "end": record.get("end"),
                "task": record.get("task"),
                "minutes": record.get("minutes", 0),
                "note": record.get("note", ""),
                "status": record.get("status", ""),
                "completion": record.get("completion"),
                "includeInStats": bool(record.get("includeInStats")),
            }
        )

    for item in milestones:
        day = by_day.setdefault(
            item["date"],
            {
                "date": item["date"],
                "minutes": 0,
                "hours": 0,
                "tasks": {},
                "notes": [],
                "records": [],
                "milestones": [],
                "taskMinutes": {},
            },
        )
        day["milestones"].append(item)

    months: list[dict] = []
    for month_start in _month_span(min(starts), max(ends)):
        next_month = month_start.replace(
            year=month_start.year + (1 if month_start.month == 12 else 0),
            month=1 if month_start.month == 12 else month_start.month + 1,
        )
        days_in_month = (next_month - timedelta(days=1)).day
        days: list[dict] = []
        for day_number in range(1, days_in_month + 1):
            date = month_start.replace(day=day_number).date().isoformat()
            day = by_day.get(
                date,
                {
                    "date": date,
                    "minutes": 0,
                    "hours": 0,
                    "tasks": {},
                    "notes": [],
                    "records": [],
                    "milestones": [],
                    "taskMinutes": {},
                },
            )
            task_items = [
                {"task": task, "minutes": minutes, "hours": round(minutes / 60, 2)}
                for task, minutes in sorted(day["tasks"].items(), key=lambda item: (-item[1], item[0]))
            ]
            days.append(
                {
                    "date": date,
                    "day": day_number,
                    "minutes": int(round(day["minutes"])),
                    "hours": round(day["minutes"] / 60, 2),
                    "tasks": task_items[:4],
                    "taskMinutes": task_items,
                    "notes": day["notes"][:4],
                    "records": sorted(day["records"], key=lambda item: item.get("start") or ""),
                    "milestones": day["milestones"][:4],
                }
            )
        months.append(
            {
                "month": month_start.strftime("%Y-%m"),
                "label": f"{month_start.year}年{month_start.month}月",
                "daysInMonth": days_in_month,
                "weekdayStart": month_start.isoweekday(),
                "days": days,
            }
        )
    return months


def _build_routine_stats(records: list[dict], milestones: list[dict] | None = None) -> dict:
    milestones = milestones or []
    included = [record for record in records if record.get("includeInStats")]
    if not included:
        return {
            "records": {"total": len(records), "included": 0},
            "range": None,
            "kpis": {"totalMinutes": 0, "totalHours": 0, "activeDays": 0, "avgMinutesPerActiveDay": 0},
            "daily": [],
            "weekly": [],
            "monthly": [],
            "tasks": [],
            "hours": [{"hour": f"{hour:02d}:00", "minutes": 0} for hour in range(24)],
            "calendar": _build_routine_calendar(records, milestones),
            "milestones": milestones,
            "notes": [],
            "recent": list(reversed(records[-20:])),
        }

    daily: dict[str, float] = {}
    weekly: dict[str, float] = {}
    monthly: dict[str, float] = {}
    tasks: dict[str, int] = {}
    hours: dict[str, float] = {f"{hour:02d}:00": 0 for hour in range(24)}
    starts: list[datetime] = []
    ends: list[datetime] = []

    for record in included:
        start = datetime.fromisoformat(record["start"])
        end = datetime.fromisoformat(record["end"])
        minutes = int(record["minutes"])
        starts.append(start)
        ends.append(end)
        tasks[record["task"]] = tasks.get(record["task"], 0) + minutes
        for key, value in _iter_time_segments(start, end, minutes, "day"):
            daily[key] = daily.get(key, 0) + value
        for key, value in _iter_time_segments(start, end, minutes, "hour"):
            hours[key] = hours.get(key, 0) + value

    start_dt = min(starts)
    end_dt = max(ends)
    for day in _date_span(start_dt, end_dt):
        day_minutes = daily.get(day, 0)
        parsed = datetime.fromisoformat(day)
        iso = parsed.isocalendar()
        week_key = f"{iso.year}-W{iso.week:02d}"
        month_key = day[:7]
        weekly[week_key] = weekly.get(week_key, 0) + day_minutes
        monthly[month_key] = monthly.get(month_key, 0) + day_minutes

    daily_series = [
        {"date": day, "minutes": int(round(daily.get(day, 0))), "hours": round(daily.get(day, 0) / 60, 2)}
        for day in _date_span(start_dt, end_dt)
    ]
    active_days = [item for item in daily_series if item["minutes"] > 0]
    active_set = {item["date"] for item in active_days}

    max_streak = 0
    streak = 0
    for item in daily_series:
        if item["minutes"] > 0:
            streak += 1
            max_streak = max(max_streak, streak)
        else:
            streak = 0
    current_streak = 0
    cursor = end_dt.date()
    while cursor.isoformat() in active_set:
        current_streak += 1
        cursor -= timedelta(days=1)

    total_minutes = int(sum(record["minutes"] for record in included))
    notes = [
        {
            "date": record["date"],
            "task": record["task"],
            "minutes": record["minutes"],
            "note": record["note"],
        }
        for record in included
        if record.get("note")
    ]
    notes.sort(key=lambda item: item["date"], reverse=True)

    return {
        "records": {"total": len(records), "included": len(included)},
        "range": {
            "start": start_dt.date().isoformat(),
            "end": end_dt.date().isoformat(),
            "firstSession": start_dt.isoformat(timespec="minutes"),
            "lastSession": end_dt.isoformat(timespec="minutes"),
        },
        "kpis": {
            "totalMinutes": total_minutes,
            "totalHours": round(total_minutes / 60, 2),
            "activeDays": len(active_days),
            "avgMinutesPerActiveDay": round(total_minutes / len(active_days)) if active_days else 0,
            "currentStreakDays": current_streak,
            "maxStreakDays": max_streak,
        },
        "daily": daily_series,
        "weekly": [
            {"week": key, "minutes": int(round(value)), "hours": round(value / 60, 2)}
            for key, value in sorted(weekly.items())
        ],
        "monthly": [
            {"month": key, "minutes": int(round(value)), "hours": round(value / 60, 2)}
            for key, value in sorted(monthly.items())
        ],
        "tasks": [
            {"task": key, "minutes": value, "hours": round(value / 60, 2)}
            for key, value in sorted(tasks.items(), key=lambda item: (-item[1], item[0]))
        ],
        "hours": [
            {"hour": key, "minutes": int(round(value)), "hours": round(value / 60, 2)}
            for key, value in sorted(hours.items())
        ],
        "calendar": _build_routine_calendar(records, milestones),
        "milestones": milestones,
        "notes": notes[:30],
        "recent": list(reversed(sorted(records, key=lambda item: item["start"])[-30:])),
    }


def _routine_payload() -> dict:
    store = _load_routine_store()
    records = store.get("records", [])
    milestones, milestone_source = _milestones_for_payload(store, records)
    return {
        "store": {
            "updatedAt": store.get("updatedAt"),
            "imports": store.get("imports", []),
            "path": str(ROUTINE_RECORDS_FILE),
            "uploadsPath": str(ROUTINE_IMPORTS_ROOT),
            "milestonesGeneratedAt": store.get("milestonesGeneratedAt"),
            "milestonesGeneratedBy": milestone_source,
        },
        "stats": _build_routine_stats(records, milestones),
        "records": records,
    }


def _import_routine_file(path: Path, original_name: str | None = None) -> dict:
    rows = _read_routine_table(path)
    parsed = _parse_routine_records(rows, original_name or path.name)
    import_info = {
        "file": original_name or path.name,
        "storedAs": str(path),
        "importedAt": datetime.now().isoformat(timespec="seconds"),
    }
    merge = _merge_routine_records(parsed["records"], import_info)
    return {"header": parsed["header"], "merge": merge, "stats": _routine_payload()["stats"]}


async def _save_routine_upload(file: UploadFile) -> Path:
    ROUTINE_IMPORTS_ROOT.mkdir(parents=True, exist_ok=True)
    suffix = Path(file.filename or "").suffix.lower()
    safe_name = _safe_filename(file.filename or f"tomatodo{suffix}")
    stored_name = f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{safe_name}"
    target = (ROUTINE_IMPORTS_ROOT / stored_name).resolve()
    if not _is_relative_to(target, ROUTINE_IMPORTS_ROOT.resolve()):
        raise HTTPException(status_code=400, detail="invalid filename")
    with target.open("wb") as handle:
        while chunk := await file.read(1024 * 1024):
            handle.write(chunk)
            if handle.tell() > 20 * 1024 * 1024:
                handle.close()
                target.unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="文件超过 20 MB")
    return target


async def _call_routine_llm(stats: dict, records: list[dict]) -> str:
    import httpx

    api_key = os.environ.get("YUSU_ROUTINE_LLM_API_KEY") or os.environ.get("DEEPSEEK_API_KEY") or os.environ.get("LLM_DEFAULT_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="未配置 DeepSeek/OpenAI-compatible API key")
    base_url = (
        os.environ.get("YUSU_ROUTINE_LLM_BASE_URL")
        or os.environ.get("DEEPSEEK_BASE_URL")
        or os.environ.get("LLM_DEFAULT_BASE_URL")
        or "https://api.deepseek.com"
    ).rstrip("/")
    model = os.environ.get("YUSU_ROUTINE_LLM_MODEL") or os.environ.get("LLM_DEFAULT_MODEL") or "deepseek-chat"
    endpoint = f"{base_url}/chat/completions"
    recent = records[-20:]
    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": "你是考研学习复盘助手。只基于给定番茄 ToDo 统计和备注，输出简洁中文建议，不编造课程进度。",
            },
            {
                "role": "user",
                "content": json.dumps(
                    {
                        "stats": stats,
                        "recentRecords": recent,
                        "要求": "总结本阶段学习节奏、薄弱信号、最高效时段、下一步建议。控制在 5 条以内。",
                    },
                    ensure_ascii=False,
                ),
            },
        ],
        "temperature": 0.2,
        "stream": False,
    }
    async with httpx.AsyncClient(timeout=45) as client:
        response = await client.post(endpoint, headers={"Authorization": f"Bearer {api_key}"}, json=payload)
    if response.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"LLM 请求失败：HTTP {response.status_code}")
    data = response.json()
    return data.get("choices", [{}])[0].get("message", {}).get("content", "").strip()


def _extract_json_array(text: str) -> list[Any]:
    clean = text.strip()
    if clean.startswith("```"):
        clean = re.sub(r"^```(?:json)?\s*", "", clean)
        clean = re.sub(r"\s*```$", "", clean)
    start = clean.find("[")
    end = clean.rfind("]")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("LLM did not return a JSON array")
    parsed = json.loads(clean[start : end + 1])
    if not isinstance(parsed, list):
        raise ValueError("LLM JSON root is not an array")
    return parsed


async def _call_routine_milestone_llm(stats: dict, records: list[dict]) -> list[dict]:
    import httpx

    note_records = [
        {
            "recordId": record.get("id"),
            "date": record.get("date"),
            "start": record.get("start"),
            "task": record.get("task"),
            "minutes": record.get("minutes"),
            "status": record.get("status"),
            "note": record.get("note"),
        }
        for record in records
        if record.get("note")
    ]
    if not note_records:
        return []

    api_key = os.environ.get("YUSU_ROUTINE_LLM_API_KEY") or os.environ.get("DEEPSEEK_API_KEY") or os.environ.get("LLM_DEFAULT_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="未配置 DeepSeek/OpenAI-compatible API key")
    base_url = (
        os.environ.get("YUSU_ROUTINE_LLM_BASE_URL")
        or os.environ.get("DEEPSEEK_BASE_URL")
        or os.environ.get("LLM_DEFAULT_BASE_URL")
        or "https://api.deepseek.com"
    ).rstrip("/")
    model = os.environ.get("YUSU_ROUTINE_LLM_MODEL") or os.environ.get("LLM_DEFAULT_MODEL") or "deepseek-chat"
    endpoint = f"{base_url}/chat/completions"
    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "你是考研例程日历标注器。只根据用户番茄 ToDo 备注抽取关键学习节点，"
                    "不要编造没有写在备注里的章节或进度。必须只返回 JSON 数组，不要 Markdown。"
                ),
            },
            {
                "role": "user",
                "content": json.dumps(
                    {
                        "notes": note_records,
                        "daily": stats.get("daily", []),
                        "allowedTypes": [
                            "start",
                            "finish",
                            "chapter",
                            "exercise",
                            "accuracy",
                            "practice",
                            "review",
                            "warning",
                            "note",
                            "other",
                        ],
                        "schema": {
                            "date": "YYYY-MM-DD，必须来自 notes",
                            "recordId": "对应备注 recordId",
                            "task": "科目或待办名称",
                            "type": "allowedTypes 之一",
                            "label": "贴在月历格子里的短标签，12 字以内最好",
                            "detail": "解释为什么这样标注，必须来自备注原文",
                            "metric": "可选对象，例如 {accuracy:'56.25%', questions:'16道'}",
                            "pinned": "是否应作为图钉显示，关键开始/完成/章节/准确率为 true",
                            "confidence": "0 到 1",
                        },
                        "examples": [
                            {
                                "note": "今天正式开始复习高数，开始零基础篇",
                                "type": "start",
                                "label": "开始高数",
                            },
                            {
                                "note": "今天正式看完了第一讲 函数的极限与连续",
                                "type": "chapter",
                                "label": "看完第一讲",
                            },
                            {
                                "note": "第一讲习题，16道题，正确率56.25%",
                                "type": "accuracy",
                                "label": "正确率56.25%",
                                "metric": {"questions": "16道", "accuracy": "56.25%"},
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
            },
        ],
        "temperature": 0.1,
        "stream": False,
    }
    async with httpx.AsyncClient(timeout=45) as client:
        response = await client.post(endpoint, headers={"Authorization": f"Bearer {api_key}"}, json=payload)
    if response.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"LLM 请求失败：HTTP {response.status_code}")
    data = response.json()
    content = data.get("choices", [{}])[0].get("message", {}).get("content", "").strip()
    try:
        return _validate_routine_milestones(_extract_json_array(content), records)
    except (json.JSONDecodeError, ValueError) as exc:
        raise HTTPException(status_code=502, detail=f"LLM 标注结果不是可用 JSON：{exc}") from exc


@app.get("/api/showcase", tags=["yusu"])
async def showcase() -> dict:
    return json.loads(DATA_FILE.read_text(encoding="utf-8"))


@app.get("/api/status", tags=["yusu"])
async def site_status() -> dict:
    projects_root = ROOT / "01_Projects"
    return {
        "vaultRoot": str(ROOT),
        "markdownFiles": len(_iter_markdown_files()),
        "projectDirectories": len([path for path in projects_root.iterdir() if path.is_dir()]),
        "rawMediaFiles": len([path for path in RAW_MEDIA_ROOT.rglob("*") if path.is_file()])
        if RAW_MEDIA_ROOT.exists()
        else 0,
        "searchMode": "live markdown scan",
        "marginaliaMode": "same-process FastAPI + source-integrated React UI/backend",
        "kaoyanDashboardAvailable": KAOYAN_DASHBOARD.is_file(),
        "routineRecords": len(_load_routine_store().get("records", [])),
    }


@app.get("/api/marginalia/status", tags=["yusu"])
async def marginalia_status() -> dict:
    settings = get_settings()
    chat = resolve_profile(settings, "chat")
    return {
        "online": True,
        "apiBase": "/v1",
        "uiBase": "/marginalia",
        "integration": "same-process",
        "backendSource": str(MARGINALIA_BACKEND),
        "semanticRecall": settings.semantic_recall_enabled,
        "semanticConfigured": bool(
            settings.semantic_recall_enabled and settings.embedding_api_key
        ),
        "embeddingModel": settings.embedding_model,
        "llmProvider": chat.provider,
        "llmModel": chat.model,
        "llmBaseUrl": chat.base_url,
        "llmKeySet": bool(chat.api_key),
        "workerEnabled": settings.worker_enabled,
    }


@app.get("/api/search", tags=["yusu"])
async def search(q: str = Query(default="")) -> dict:
    return {"query": q, "results": search_markdown(q)}


@app.get("/api/doc", tags=["yusu"])
async def document(path: str = Query(default="")) -> dict:
    target = _resolve_markdown_doc(path)
    if target is None:
        raise HTTPException(status_code=404, detail="markdown document not found")
    text = _read_text(target)
    return {
        "title": _first_heading(text, target.stem),
        "path": target.relative_to(ROOT).as_posix(),
        "content": text,
        "lines": len(text.splitlines()),
    }


@app.get("/api/routine/stats", tags=["yusu"])
async def routine_stats() -> dict:
    return _routine_payload()


@app.post("/api/routine/import", tags=["yusu"])
async def routine_import(file: UploadFile = File(...)) -> dict:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".xls", ".xlsx", ".csv", ".tsv"}:
        raise HTTPException(status_code=415, detail="仅支持 .xls, .xlsx, .csv, .tsv")
    target = await _save_routine_upload(file)
    return _import_routine_file(target, file.filename)


@app.post("/api/routine/import-json", tags=["yusu"])
async def routine_import_json(
    file: UploadFile = File(...),
    records_json: str = Form(...),
) -> dict:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".xls", ".xlsx", ".csv", ".tsv"}:
        raise HTTPException(status_code=415, detail="仅支持 .xls, .xlsx, .csv, .tsv")
    try:
        raw_records = json.loads(records_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=422, detail=f"记录 JSON 解析失败：{exc}") from exc
    if not isinstance(raw_records, list):
        raise HTTPException(status_code=422, detail="记录 JSON 必须是数组")
    target = await _save_routine_upload(file)
    records = _normalize_client_routine_records(raw_records, file.filename or target.name)
    if not records:
        raise HTTPException(status_code=422, detail="浏览器解析结果里没有有效番茄记录")
    merge = _merge_routine_records(
        records,
        {
            "file": file.filename or target.name,
            "storedAs": str(target),
            "importedAt": datetime.now().isoformat(timespec="seconds"),
            "parser": "browser-sheetjs",
        },
    )
    return {"merge": merge, "stats": _routine_payload()["stats"]}


@app.post("/api/routine/summary", tags=["yusu"])
async def routine_summary() -> dict:
    payload = _routine_payload()
    records = [record for record in payload["records"] if record.get("includeInStats")]
    if not records:
        raise HTTPException(status_code=400, detail="还没有可总结的有效打卡记录")
    text = await _call_routine_llm(payload["stats"], records)
    return {"summary": text, "modelConfigured": bool(text)}


@app.post("/api/routine/milestones", tags=["yusu"])
async def routine_milestones() -> dict:
    store = _load_routine_store()
    records = store.get("records", [])
    if not any(record.get("note") for record in records):
        raise HTTPException(status_code=400, detail="还没有带备注的记录可标注")
    stats = _build_routine_stats(records, _fallback_routine_milestones(records))
    milestones = await _call_routine_milestone_llm(stats, records)
    if not milestones:
        raise HTTPException(status_code=502, detail="模型没有返回可用关键节点")
    store["milestones"] = milestones
    store["milestonesGeneratedAt"] = datetime.now().isoformat(timespec="seconds")
    store["milestonesGeneratedBy"] = "deepseek"
    store["milestoneSignature"] = _routine_note_signature(records)
    _save_routine_store(store)
    return {"milestones": milestones, "payload": _routine_payload()}


@app.get("/media/raw/{filename:path}", include_in_schema=False)
async def raw_media(filename: str) -> FileResponse:
    target = (RAW_MEDIA_ROOT / unquote(filename)).resolve()
    if not _is_relative_to(target, RAW_MEDIA_ROOT.resolve()) or not target.is_file():
        raise HTTPException(status_code=404, detail="media file not found")
    return FileResponse(target, headers={"Cache-Control": "public, max-age=3600"})


@app.get("/marginalia", include_in_schema=False)
async def marginalia_root() -> RedirectResponse:
    return RedirectResponse(url="/marginalia/chat", status_code=307)


@app.get("/marginalia/{asset_path:path}", include_in_schema=False)
async def marginalia_ui(asset_path: str) -> FileResponse:
    clean = unquote(asset_path).replace("\\", "/").lstrip("/")
    target = (MARGINALIA_DIST / clean).resolve()
    if _is_relative_to(target, MARGINALIA_DIST.resolve()) and target.is_file():
        return FileResponse(target)
    index = MARGINALIA_DIST / "index.html"
    if not index.is_file():
        raise HTTPException(
            status_code=503,
            detail="Marginalia UI is not built. Run tools/build-yusu-integrated-marginalia-ui.ps1",
        )
    return FileResponse(index, headers={"Cache-Control": "no-store, max-age=0"})


@app.get("/api/kaoyan/status", tags=["yusu"])
async def kaoyan_status() -> dict:
    return {
        "online": KAOYAN_DASHBOARD.is_file(),
        "integration": "same-process static dashboard from source workspace",
        "workspace": str(KAOYAN_WORKSPACE),
        "dashboard": str(KAOYAN_DASHBOARD),
        "dashboardUpdated": _file_timestamp(KAOYAN_DASHBOARD),
        "dashboardBytes": KAOYAN_DASHBOARD.stat().st_size if KAOYAN_DASHBOARD.is_file() else 0,
        "uiBase": "/kaoyan/",
        "privacyBoundary": "generated dashboard data stays in the exam project workspace and is not copied into this vault",
    }


@app.get("/kaoyan", include_in_schema=False)
async def kaoyan_root_redirect() -> RedirectResponse:
    return RedirectResponse(url="/kaoyan/", status_code=307)


@app.get("/kaoyan/{asset_path:path}", include_in_schema=False)
async def kaoyan_dashboard(asset_path: str = ""):
    target = _resolve_kaoyan_asset(asset_path)
    if target is None:
        detail = (
            "Kaoyan dashboard not found. Set YUSU_KAOYAN_WORKSPACE to the exam project root."
            if not KAOYAN_DASHBOARD.is_file()
            else "Kaoyan asset is not available or not allowed."
        )
        raise HTTPException(status_code=404, detail=detail)
    if target.resolve() == KAOYAN_DASHBOARD.resolve():
        html = target.read_text(encoding="utf-8-sig", errors="replace")
        return HTMLResponse(
            _inject_kaoyan_return_button(html),
            headers={"Cache-Control": "no-store, max-age=0"},
        )
    return FileResponse(target, headers={"Cache-Control": "no-store, max-age=0"})


@app.get("/routine", include_in_schema=False)
async def routine_root_redirect() -> RedirectResponse:
    return RedirectResponse(url="/routine/", status_code=307)


@app.get("/routine/", include_in_schema=False)
async def routine_page() -> FileResponse:
    return FileResponse(
        WEB_ROOT / "routine.html",
        headers={"Cache-Control": "no-store, max-age=0"},
    )


@app.get("/routine.css", include_in_schema=False)
async def routine_styles() -> FileResponse:
    return FileResponse(WEB_ROOT / "routine.css", media_type="text/css")


@app.get("/routine.js", include_in_schema=False)
async def routine_javascript() -> FileResponse:
    return FileResponse(WEB_ROOT / "routine.js", media_type="text/javascript")


@app.get("/styles.css", include_in_schema=False)
async def site_styles() -> FileResponse:
    return FileResponse(WEB_ROOT / "styles.css", media_type="text/css")


@app.get("/app.js", include_in_schema=False)
async def site_javascript() -> FileResponse:
    return FileResponse(WEB_ROOT / "app.js", media_type="text/javascript")


@app.get("/", include_in_schema=False)
async def site_root() -> FileResponse:
    return FileResponse(
        WEB_ROOT / "index.html",
        headers={"Cache-Control": "no-store, max-age=0"},
    )


app.mount("/assets", StaticFiles(directory=WEB_ROOT / "assets"), name="yusu-assets")
app.mount("/vendor", StaticFiles(directory=WEB_ROOT / "vendor"), name="yusu-vendor")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the integrated YUSU + Marginalia site.")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8787)
    args = parser.parse_args()

    import uvicorn

    uvicorn.run(app, host=args.host, port=args.port, log_level="info")


if __name__ == "__main__":
    main()
